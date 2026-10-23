"""Gera o DOCX do artigo no layout do modelo PLN2026, sem editar origens.

Executar com o Python das dependências de documentos do Codex, com python-docx.
O PDF é exportado pelo Word e conferido separadamente; não é PDF de outro texto.
"""
from pathlib import Path
from datetime import datetime, timezone
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

BASE = Path(__file__).resolve().parent


def font(run, size=10, bold=False, name='Times New Roman'):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    if fonts is not None:
        fonts.set(qn('w:ascii'), name)
        fonts.set(qn('w:hAnsi'), name)
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'pt-BR')
    props.append(language)


def inline(paragraph, text, size=10, bold=False, name='Times New Roman'):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|ᵈᵉ)'
    for part in re.split(pattern, text):
        if not part:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', part)
        if link:
            rel = paragraph.part.relate_to(link[2], 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
            element = OxmlElement('w:hyperlink')
            element.set(qn('r:id'), rel)
            run = paragraph.add_run(link[1])
            font(run, size, bold, name)
            paragraph._p.remove(run._r)
            element.append(run._r)
            paragraph._p.append(element)
        else:
            strong = part.startswith('**') and part.endswith('**')
            run = paragraph.add_run(part[2:-2] if strong else part)
            font(run, size, bold or strong, 'Cambria' if part == 'ᵈᵉ' else name)


def paragraph(doc, text, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, after=0, name='Times New Roman'):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Mm(5) if indent else Mm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(12 if size == 10 else size + 2)
    pf.widow_control = True
    inline(p, text, size, bold, name)
    return p


def columns(section, count):
    element = section._sectPr.find(qn('w:cols'))
    if element is None:
        element = OxmlElement('w:cols')
        section._sectPr.append(element)
    element.set(qn('w:num'), str(count))
    element.set(qn('w:space'), '567')


def table(doc, rows):
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [Mm(60), Mm(15), Mm(15)]
    for i, width in enumerate(widths):
        t.columns[i].width = width
    for index, row in enumerate(rows):
        cells = t.rows[0].cells if index == 0 else t.add_row().cells
        props = t.rows[index]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        props.append(no_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            props.append(repeat)
        for i, text in enumerate(row):
            c = cells[i]
            c.width = widths[i]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.keep_with_next = True
            inline(p, text.strip(), bold=index == 0)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
        b = OxmlElement('w:' + edge)
        b.set(qn('w:val'), 'single' if edge in ('top', 'bottom') else 'nil')
        b.set(qn('w:sz'), '6')
        b.set(qn('w:color'), '000000')
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for c in t.rows[0].cells:
        b = OxmlElement('w:tcBorders')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        b.append(bottom)
        c._tc.get_or_add_tcPr().append(b)
    return t


def main():
    source = (BASE / 'artigo.md').read_text(encoding='utf-8')
    blocks = source.strip().split('\n\n')
    doc = Document()
    first = doc.sections[0]
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(13)
        section.bottom_margin = Mm(13)
        section.left_margin = Mm(10)
        section.right_margin = Mm(10)
        section.header_distance = Mm(5)
        section.footer_distance = Mm(6)
    columns(first, 1)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    # O template padrão do python-docx contém bordas de tema nos estilos de título.
    # O modelo PLN2026 não tem régua decorativa sob o título nem sob seções.
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for style_name in ['Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[style_name]
        style.font.name = 'Arial'
        style.font.size = Pt(14 if style_name == 'Title' else 12)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if style_name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.first_line_indent = Mm(0)
    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline(title, blocks[0][2:], 14, True, 'Arial')
    paragraph(doc, blocks[1], size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=4)
    paragraph(doc, blocks[2], align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6)
    paragraph(doc, blocks[3], align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    paragraph(doc, blocks[4], align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=8)
    in_body = False
    in_references = False
    i = 5
    while i < len(blocks):
        block = blocks[i]
        if block.startswith('## Introdução'):
            doc.add_section(WD_SECTION_START.CONTINUOUS)
            columns(doc.sections[-1], 2)
            in_body = True
        if block.startswith('## ') or block.startswith('### '):
            heading = block.lstrip('# ').strip()
            if not in_body:
                next_block = blocks[i + 1]
                paragraph(doc, f'**{heading}:** {next_block}', indent=False, after=4)
                i += 2
                continue
            in_references = heading == 'Referências bibliográficas'
            p = doc.add_paragraph(style='Heading 2' if block.startswith('### ') else 'Heading 1')
            inline(p, heading, 12, True, 'Arial')
        elif block.startswith('|'):
            data = [line.strip().strip('|').split('|') for line in block.splitlines()]
            table(doc, [data[0], *data[2:]])
        elif block.startswith('**Tabela'):
            p = paragraph(doc, block, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=4)
            p.paragraph_format.keep_with_next = True
        elif in_references:
            p = paragraph(doc, block, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, after=5)
            p.paragraph_format.left_indent = Mm(5)
            p.paragraph_format.first_line_indent = Mm(-5)
        else:
            is_keywords = block.startswith(('**Palavras-chave:', '**Keywords:'))
            paragraph(doc, block, indent=in_body and not is_keywords, after=6 if is_keywords else 0)
        i += 1
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not footer.text and not footer._p.xpath('.//w:fldSimple'):
            field = OxmlElement('w:fldSimple')
            field.set(qn('w:instr'), 'PAGE \\# "000"')
            run = OxmlElement('w:r')
            text = OxmlElement('w:t')
            text.text = '001'
            run.append(text)
            field.append(run)
            footer._p.append(field)
            font(footer.add_run(' | Página'), 8)
    doc.core_properties.title = blocks[0][2:]
    doc.core_properties.author = blocks[1]
    doc.core_properties.subject = 'PI IV e PLN — diagnóstico de desenvolvimento do PaleographIA'
    doc.core_properties.keywords = 'HTR, português histórico, CER, WER, referência assistida'
    doc.core_properties.comments = 'Artigo do projeto PaleographIA; fonte textual em artigo.md.'
    generated_at = datetime.now(timezone.utc)
    doc.core_properties.created = generated_at
    doc.core_properties.modified = generated_at
    doc.save(BASE / 'artigo.docx')
    print('Gerado:', BASE / 'artigo.docx')


if __name__ == '__main__':
    main()
